"""Seguimiento posterior al corte (campos de validación v2): misma semana a ambos lados.

Origen: en la tabla W37 `pron_sem_previa` era el pronóstico de la W38 y `realidad_sem_previa` la
realidad de la W39. Estas pruebas fijan que el par se forma por (serie, semana del boletín, ds),
que nada se desplaza y que un lado faltante declara el par ausente.
"""

from __future__ import annotations

import importlib
from pathlib import Path
import sys
from typing import Any

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest
from scripts import genera_validacion_semanal as gvs

from epiforecast.utils.semana_epi import ds_de_semana_boletin

pytestmark = pytest.mark.unit


@pytest.fixture
def seg(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Any:
    """El módulo, importado con ``conf["paths"]["reports"]`` apuntando a un directorio temporal."""
    cfg = sys.modules["epiforecast.utils.config"]
    monkeypatch.setitem(cfg.conf["paths"], "reports", str(tmp_path / "reports"))
    for m in ("scripts.genera_seguimiento_post_corte", "scripts.build_dengue_gallery"):
        monkeypatch.delitem(sys.modules, m, raising=False)
    return importlib.import_module("scripts.genera_seguimiento_post_corte")


CORTE = (2026, 37)
SERIE = {
    "cohorte": "neuro",
    "padecimiento": "Depresion",
    "entidad": "Nacional",
    "sexo": "general",
    "motor": "Prophet",
}


def test_semanas_objetivo_cruzan_el_anio_con_la_53(seg: Any) -> None:
    assert seg.semanas_objetivo((2025, 51), (2026, 2)) == [
        (2025, 52),
        (2025, 53),
        (2026, 1),
        (2026, 2),
    ]
    assert seg.semanas_objetivo((2026, 51), (2027, 1)) == [(2026, 52), (2027, 1)]  # 2026 tiene 52
    assert seg.semanas_objetivo(CORTE, (2026, 39)) == [(2026, 38), (2026, 39)]
    assert seg.semanas_objetivo(CORTE, CORTE) == []


def test_el_paso_sin_semana_del_boletin_se_declara_fuera(seg: Any) -> None:
    fc = pd.DataFrame({"ds": pd.to_datetime(["2026-12-21", "2026-12-28", "2027-01-04"])})
    fuera = seg.pasos_sin_semana(fc, pd.Timestamp("2026-12-14"), pd.Timestamp("2027-01-04"))
    assert fuera == ["2026-12-28"]


def _fc(filas: list[tuple[int, int, float]], corte: tuple[int, int] = CORTE) -> pd.DataFrame:
    """Pronóstico sintético de la SERIE: (año, semana del boletín, yhat), con su paso."""
    out = []
    for i, (a, s, y) in enumerate(filas, start=1):
        out.append(
            {
                "ds": ds_de_semana_boletin(a, s),
                "yhat": y,
                "meta_padecimiento": "Depresión",
                "meta_entidad": None,  # el nacional viaja sin entidad
                "meta_modo": "general",
                "anio_boletin_corte": corte[0],
                "semana_boletin_corte": corte[1],
                "fase_modelo": "pronostico",
                "adelanto_semanas": i,
            }
        )
    return pd.DataFrame(out)


class _Obs:
    """Observado sintético: {(año, semana): valor}; ausente si no está."""

    def __init__(self, valores: dict[tuple[int, int], float], ultima: tuple[int, int]) -> None:
        self.valores, self.ultima = valores, ultima

    def ultima_semana(self, cohorte: str) -> tuple[int, int]:
        return self.ultima

    def valor(self, fila, anio, semana, ds):  # noqa: ANN001, ANN201
        v = self.valores.get((anio, semana))
        return (v, "") if v is not None else (None, "observación no desagregable")


def _corre(seg: Any, monkeypatch, fc: pd.DataFrame, obs: _Obs) -> pd.DataFrame:  # noqa: ANN001
    monkeypatch.setattr(seg, "asignacion", lambda: [dict(SERIE)])
    monkeypatch.setattr(seg, "Observado", lambda: obs)
    monkeypatch.setattr(seg, "_load_fc", lambda motor: fc)
    monkeypatch.setattr(seg, "_sha", lambda p: "x")
    tabla, resumen = seg.construye(CORTE)
    assert resumen["semanas_evaluadas"] == len(tabla)
    return tabla.set_index("semana_objetivo")


def test_cada_semana_se_empareja_consigo_misma(seg: Any, monkeypatch) -> None:  # noqa: ANN001
    t = _corre(
        seg,
        monkeypatch,
        _fc([(2026, 38, 100.0), (2026, 39, 200.0)]),
        _Obs({(2026, 38): 110, (2026, 39): 190}, (2026, 39)),
    )
    assert t.loc["2026-W38", ["pronostico", "real", "horizonte", "corte"]].tolist() == [
        100,
        110,
        1,
        "2026-W37",
    ]
    assert t.loc["2026-W39", ["pronostico", "real", "horizonte"]].tolist() == [200, 190, 2]
    assert t.loc["2026-W39", "ds_objetivo"] == "2026-09-14"
    assert (t["estado"] == "par").all()


def test_sin_pronostico_de_la_semana_no_se_usa_el_de_la_anterior(seg: Any, monkeypatch) -> None:  # noqa: ANN001
    """El error de la tabla W37: realidad W39 junto al pronóstico W38. Aquí debe quedar ausente."""
    t = _corre(
        seg,
        monkeypatch,
        _fc([(2026, 38, 100.0)]),
        _Obs({(2026, 38): 110, (2026, 39): 190}, (2026, 39)),
    )
    w39 = t.loc["2026-W39"]
    assert w39["estado"] == "ausente" and "sin pronóstico guardado" in w39["motivo"]
    assert pd.isna(w39["pronostico"]) and pd.isna(w39["error_abs"])
    assert w39["real"] == 190  # lo observado se conserva, sin pareja


def test_sin_observacion_el_par_queda_ausente(seg: Any, monkeypatch) -> None:  # noqa: ANN001
    t = _corre(
        seg,
        monkeypatch,
        _fc([(2026, 38, 100.0), (2026, 39, 200.0)]),
        _Obs({(2026, 39): 190}, (2026, 39)),
    )
    assert t.loc["2026-W38", "estado"] == "ausente"
    assert pd.isna(t.loc["2026-W38", "real"]) and t.loc["2026-W38", "pronostico"] == 100


def test_un_pronostico_de_otro_corte_no_entra(seg: Any, monkeypatch) -> None:  # noqa: ANN001
    t = _corre(
        seg,
        monkeypatch,
        _fc([(2026, 38, 100.0), (2026, 39, 200.0)], corte=(2026, 33)),
        _Obs({(2026, 38): 1, (2026, 39): 1}, (2026, 39)),
    )
    assert set(t["motivo"]) == {"pronóstico de otro corte"}


def test_paso_del_modelo_distinto_del_horizonte_no_se_empareja(seg: Any, monkeypatch) -> None:  # noqa: ANN001
    fc = _fc([(2026, 38, 100.0), (2026, 39, 200.0)])
    fc.loc[fc["ds"] == ds_de_semana_boletin(2026, 39), "adelanto_semanas"] = 3
    t = _corre(seg, monkeypatch, fc, _Obs({(2026, 38): 1, (2026, 39): 1}, (2026, 39)))
    assert "paso 3 del pronóstico ≠ horizonte 2" in t.loc["2026-W39", "motivo"]


def test_la_tabla_lleva_ambos_lados_de_la_misma_semana(tmp_path: Path) -> None:
    ruta = tmp_path / "tabla.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Produccion"
    ws.append(["padecimiento", "entidad", "sexo", "pron_sem_previa", "realidad_sem_previa"])
    ws.append(["Depresion", "Nacional", "general", 999, None])  # 999: el pronóstico de la W38
    ws.append(["Depresion", "Ciudad de México", "mujeres", 999, None])
    ws.append(["Dengue", "Nacional", "general", None, None])
    wb.save(ruta)
    base = {
        "cohorte": "neuro",
        "padecimiento": "Depresion",
        "corte": "2026-W37",
        "semana_objetivo": "2026-W39",
        "horizonte": 2,
    }
    ultima = pd.DataFrame(
        [
            {
                **base,
                "entidad": "Nacional",
                "sexo": "general",
                "pronostico": 3136,
                "real": 3332,
                "estado": "par",
                "motivo": "",
            },
            {
                **base,
                "entidad": "Ciudad de México",
                "sexo": "mujeres",
                "pronostico": 50,
                "real": None,
                "estado": "ausente",
                "motivo": "revisión",
            },
        ]
    )

    assert gvs._update_excel(ultima, ruta) == 2

    ws = load_workbook(ruta)["Produccion"]
    filas = [[c.value for c in fila] for fila in ws.iter_rows(min_row=1)]
    assert filas[0][-4:] == list(gvs.COLUMNAS_VALIDACION)
    assert filas[1][3:] == [3136, 3332, "2026-W39", "2026-W37", 2, "par"]
    assert filas[2][3:5] == [None, None]  # ausente: ni el 999 de otra semana ni medio par
    assert filas[2][-1] == "ausente: revisión"
    assert filas[3][3:5] == [None, None] and filas[3][5] is None  # Dengue no está en esta tabla


def test_la_tabla_rechaza_mezclar_semanas(tmp_path: Path) -> None:
    dos = pd.DataFrame({"semana_objetivo": ["2026-W38", "2026-W39"], "cohorte": ["neuro"] * 2})
    with pytest.raises(ValueError, match="UNA semana"):
        gvs._update_excel(dos, tmp_path / "x.xlsx")


# ── el knowledge del EpiBot ────────────────────────────────────────────────────


def _resumen(tmp_path: Path, ultima: str = "2026-W39") -> Path:
    import json

    r = {"semanas_evaluadas": 2, "semanas": [{"etiqueta": "2026-W38"}, {"etiqueta": ultima}]}
    ruta = tmp_path / "seguimiento_post_corte.json"
    ruta.write_text(json.dumps(r), "utf-8")
    return ruta


def test_knowledge_exige_que_tabla_y_resumen_hablen_de_la_misma_semana(tmp_path: Path) -> None:
    from scripts.build_web_knowledge import carga_seguimiento

    ok = [{"semana_validacion": "2026-W39"}] * 3
    assert carga_seguimiento(_resumen(tmp_path), ok)["semanas_evaluadas"] == 2
    with pytest.raises(SystemExit, match="2026-W38"):
        carga_seguimiento(_resumen(tmp_path), [{"semana_validacion": "2026-W38"}])
    with pytest.raises(SystemExit, match="la tabla declara"):
        carga_seguimiento(_resumen(tmp_path), [{"semana_validacion": None}])


def test_tabla_v1_sin_semana_declarada_no_publica_el_par(tmp_path: Path) -> None:
    from types import SimpleNamespace

    from scripts.build_web_knowledge import build_prod_models

    v1 = pd.DataFrame(
        [
            {
                "padecimiento": "Depresion",
                "entidad": "Nacional",
                "sexo": "general",
                "modelo_produccion": "Prophet",
                "pron_sem_previa": 3118,
                "realidad_sem_previa": 3332,
            }
        ]
    )
    fila = build_prod_models(SimpleNamespace(prod_models=v1))[0]
    assert "pron_sem_previa" not in fila and "realidad_sem_previa" not in fila

    v2 = v1.assign(
        semana_validacion="2026-W39", corte_validacion="2026-W37", horizonte_validacion=2
    )
    fila = build_prod_models(SimpleNamespace(prod_models=v2))[0]
    assert (fila["pron_sem_previa"], fila["realidad_sem_previa"]) == (3118, 3332)
    assert (fila["semana_validacion"], fila["horizonte_validacion"]) == ("2026-W39", 2)
